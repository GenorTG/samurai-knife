"""Live editing of the OPEN office document through the Euro-Office plugin, against REAL docker and a real editor:
install the office sidecar (the plugin is mounted automatically), open a python-docx-made docx, the plugin connects, run
every word command through the bridge and verify the effect in the editor (reads back through the plugin, Ctrl+A
selection, Ctrl+Z undo) and in the saved file after a force-save (python-docx + the raw XML); basic xlsx and pptx
commands checked in the saved files; an assistant run with the scripted mock provider (CALL: lines) that reads and edits
the open document, saves it, and is stopped by the trust modes for destructive calls; the tools disappear when no editor
is connected; a changed plugin is served without a restart; a container from an older compose file is recreated by
the next lazy start.

Isolation: own scratch worktree server on the Tailscale address (a container cannot reach the host loopback), own
KNIFE_DATA_DIR, port and compose project (KNIFE_SIDECAR_PREFIX=sk-live-<pid>), own mock provider; removes the sidecar
(container + volumes; the image stays cached unless --remove-image) and the worktree at the end.

    python3 tests/office_live_e2e.py [--host <tailscale-or-LAN-ip>] [--hold] [--remove-image]
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import urllib.request
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from office_e2e import Api, container_names, wait_editor_ready, wait_for  # noqa: E402
from scratch_server import ROOT, ScratchServer  # noqa: E402

from playwright.sync_api import Page, sync_playwright  # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []
SHOTS = ROOT / ".tmp" / "office-live-e2e"
PLUGIN = ROOT / "sidecars" / "office" / "plugins" / "knife-bridge"


def check(name: str, ok: bool, detail: object = "") -> bool:
    text = detail if isinstance(detail, str) else json.dumps(detail, ensure_ascii=False)[:600]
    RESULTS.append((name, ok, text))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}{f' — {text}' if text and not ok else ''}", flush=True)
    return ok


def venv_python() -> str:
    explicit = os.environ.get("KNIFE_PYTHON")
    if explicit:
        return explicit
    home = os.environ.get("KNIFE_PY_HOME") or str(Path.home() / "work" / "samurai-knife-py")
    return str(Path(home) / "venv" / "bin" / "python")


def py_docx(script: str, payload: bytes = b"") -> dict:
    """Runs python-docx (project venv) on bytes from stdin; the script prints one JSON object."""
    out = subprocess.run([venv_python(), "-I", "-c", script], input=payload, capture_output=True, timeout=60)
    if out.returncode != 0:
        raise RuntimeError(out.stderr.decode()[-800:])
    return json.loads(out.stdout)


MAKE_DOCX = r"""
import io, json, sys, base64
from docx import Document
d = Document()
d.add_heading("Quarterly report", level=1)
d.add_paragraph("Draft line one about revenue.")
d.add_paragraph("Draft line two about costs.")
d.add_paragraph("Obsolete paragraph to delete.")
buf = io.BytesIO(); d.save(buf)
print(json.dumps({"b64": base64.b64encode(buf.getvalue()).decode()}))
"""

READ_DOCX = r"""
import io, json, sys
from docx import Document
d = Document(io.BytesIO(sys.stdin.buffer.read()))
print(json.dumps({"paragraphs": [{"text": p.text, "style": p.style.name} for p in d.paragraphs]}))
"""


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def download(api: Api, doc_id: str) -> bytes:
    status, raw = api.call("GET", f"/api/office/docs/{doc_id}?download=1")
    assert status == 200, status
    return raw


def force_save(api: Api, doc_id: str) -> tuple[bool, str]:
    """Server force-save (Document Server command service); retried while the editor has not sent its last change yet."""
    before = api.doc(doc_id)["version"]
    status = ""
    for _ in range(8):
        _, body = api.call("POST", f"/api/office/docs/{doc_id}", {"action": "forcesave"})
        status = body.get("status", str(body))
        if status == "saving":
            break
        time.sleep(1.2)
    if status != "saving":
        return False, status
    ok = wait_for(lambda: api.doc(doc_id)["version"] > before, timeout=60, step=0.5)
    return bool(ok), f"{status}, version {before} -> {api.doc(doc_id)['version']}"


def _get(url: str) -> str:
    try:
        return urllib.request.urlopen(url, timeout=5).read().decode().strip()
    except OSError:
        return ""


def zip_text(raw: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return z.read(name).decode()


def zip_names(raw: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return z.namelist()


def live(page: Page, cmd: str, args: dict | None = None) -> dict:
    return page.evaluate("([c, a]) => window.knifeOfficeLive.call(c, a)", [cmd, args or {}])


def wait_live(page: Page, timeout_ms: int = 90_000) -> str:
    try:
        page.wait_for_function(
            "() => /connected \\(|unavailable/.test(document.querySelector('[data-agent-region=\"office-live-status\"]')?.textContent ?? '')",
            timeout=timeout_ms,
        )
    except Exception:
        page.screenshot(path=str(SHOTS / "wait-live-timeout.png"))
        status = page.locator('[data-agent-region="office-live-status"]')
        print("live status at timeout:", status.inner_text() if status.count() else "(no status region)", flush=True)
        raise
    return page.locator('[data-agent-region="office-live-status"]').inner_text()


def open_doc(page: Page, url: str, doc_id: str) -> tuple[str, str]:
    page.goto(f"{url}/tools/office?doc={doc_id}")
    ready = wait_editor_ready(page)
    return ready, wait_live(page)


def editor_canvas(page: Page):
    frame = page.frame_locator('iframe[name="frameEditor"]')
    canvas = frame.locator("#editor_sdk")
    canvas.wait_for(timeout=60_000)
    return canvas


def close_editor(page: Page) -> None:
    page.get_by_role("button", name="Close editor").click()
    page.get_by_role("heading", name="Your documents").wait_for(timeout=30_000)


class Assistant:
    """The widget driven with the scripted mock provider (CALL: lines), as in agent_trust_modes_e2e.py."""

    def __init__(self, page: Page, mock: str) -> None:
        self.page = page
        self.mock = mock
        self.n = 0
        self.key = ""

    def open(self) -> None:
        # Keyboard focus may sit inside the editor iframe (it swallows Ctrl+J): focus the page first.
        self.page.locator("main h2").first.click()
        self.page.keyboard.press("Control+j")
        self.page.wait_for_selector("[role=dialog][aria-label='Toolbox assistant']")
        self.key = self.page.evaluate("fetch('/api/account').then(r => r.json()).then(a => `sk.agent.chat.v1.${a.id}:${encodeURIComponent(a.provider?.baseUrl ?? '')}:${encodeURIComponent(a.provider?.model ?? '')}`)")
        self.page.locator("[aria-label='Ask before acting'] [role=radio]").first.wait_for()

    def mode(self, label: str) -> None:
        self.page.locator("[aria-label='Ask before acting']").get_by_role("radio", name=label, exact=True).click()

    def run(self, text: str, decide: str = "Deny") -> str:
        """Send a prompt; 'ran', or 'asked:<title>' / 'asked!:<title>' (destructive) after answering with `decide`."""
        self.n += 1
        text = f"live run {self.n}\n{text}"
        box = self.page.locator("textarea[aria-label='Message the assistant']")
        box.fill(text)
        box.press("Enter")
        done = """([k, expected]) => { const m = JSON.parse(localStorage.getItem(k) || '[]');
            const users = m.filter(x => x.role === 'user');
            if (users.at(-1)?.content !== expected || document.documentElement.hasAttribute('data-agent-active')) return false;
            const l = m[m.length - 1]; return (l.role === 'assistant' && !l.tool_calls) || l.role === 'notice'; }"""
        asked = ""
        for _ in range(900):
            card = self.page.locator("[data-agent-approval]")
            if card.count() and not asked:
                kind = card.get_attribute("data-agent-approval")
                title = self.page.locator("[data-agent-approval] p").first.inner_text().replace("⚠ ", "")
                asked = ("asked!:" if kind == "destructive" else "asked:") + title
                self.page.locator("[data-agent-approval]").get_by_role("button", name=decide, exact=True).click()
            if self.page.evaluate(f"({done})", [self.key, text]):
                return asked or "ran"
            self.page.wait_for_timeout(100)
        return "timeout"

    def results(self, name: str) -> list[dict]:
        """Tool results of `name` in the last run (newest last)."""
        ms = json.loads(self.page.evaluate("k => localStorage.getItem(k) || '[]'", self.key))
        last_user = max(i for i, m in enumerate(ms) if m["role"] == "user")
        ids = {c["id"] for m in ms[last_user:] if m["role"] == "assistant" for c in m.get("tool_calls", []) if c["function"]["name"] == name}
        out = []
        for m in ms[last_user:]:
            if m["role"] == "tool" and m["tool_call_id"] in ids:
                try:
                    out.append(json.loads(m["content"]))
                except ValueError:
                    out.append({"raw": m["content"]})
        return out

    def sent_tools(self) -> list[str]:
        body = json.loads(urllib.request.urlopen(self.mock + "/__last", timeout=10).read())
        request = body.get("body", body)
        return [t["function"]["name"] for t in request.get("tools", [])]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--host", default=os.environ.get("KNIFE_E2E_HOST") or os.environ.get("KNIFE_HOST"))
    parser.add_argument("--hold", action="store_true", help="before cleanup, keep everything up until Enter")
    parser.add_argument("--remove-image", action="store_true")
    args = parser.parse_args()

    SHOTS.mkdir(parents=True, exist_ok=True)
    prefix = f"sk-live-{os.getpid()}"
    project = f"{prefix}-office"
    mock_port = free_port()
    mock = subprocess.Popen(["node", str(ROOT / "tests" / "mock_provider.mjs"), str(mock_port)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    mock_url = f"http://127.0.0.1:{mock_port}"
    server = ScratchServer(host=args.host)
    # Untracked files (a fresh plugin, new modules) are not in the stash-based worktree: overlay them.
    url = server.start(env={"KNIFE_SIDECAR_PREFIX": prefix}, overlay=["src", "sidecars"])
    api = Api(url)
    print(f"scratch server {url}, data dir {server.data_dir}, compose project {project}, mock {mock_url}", flush=True)
    try:
        status, _ = api.call("PUT", "/api/account/provider", {"baseUrl": f"{mock_url}/v1", "model": "mock-model", "apiKey": "mock-key"})
        check("mock provider configured", status == 200, str(status))

        # ------------------------------------------------------------------ install: the plugin comes with the sidecar
        status, body = api.call("POST", "/api/services/office", {"action": "install"})
        op = api.wait_op(body["op"], timeout=3600)
        if not check("install op done", op["status"] == "done", op.get("error") or op["phase"]):
            print("\n".join(op["log"][-30:]))
            return 1
        check("install log names the synced editor plugin", any("editor plugins updated: knife-bridge-" in line for line in op["log"]), op["log"][:6])
        _, body = api.call("GET", "/api/services/office")
        svc = body["service"]
        public = svc["publicUrl"]
        plugins = json.loads(urllib.request.urlopen(public + "/plugins.json", timeout=10).read())
        folders = [p.split("/")[-2] for p in plugins.get("pluginsData", [])]
        check("the Document Server lists the plugin in /plugins.json", len(folders) == 1 and folders[0].startswith("knife-bridge-"), plugins)
        mounted = Path(server.data_dir) / "services" / "office" / "assets" / folders[0]
        check("mounted copy equals the repo plugin", all((mounted / f.name).read_bytes() == f.read_bytes() for f in PLUGIN.iterdir()), str(mounted))
        served = urllib.request.urlopen(f"{public}/sdkjs-plugins/{folders[0]}/config.json", timeout=10).read()
        check("nginx serves the plugin files", json.loads(served)["guid"] == json.loads((PLUGIN / "config.json").read_text())["guid"])
        check("status: no outdated-container warning right after install", not any("older settings" in w for w in svc["warnings"]), svc["warnings"])

        # ------------------------------------------------------------------ editor config: origin rule
        made = py_docx(MAKE_DOCX)
        status, body = api.call("POST", "/api/office/docs", {"action": "upload", "name": "Live report.docx", "data": made["b64"]})
        check("python-docx document uploaded", status == 201, str(status))
        docx = body["doc"]
        status, body = api.call("GET", f"/api/office/docs/{docx['id']}/editor?origin={url}")
        guid = json.loads((PLUGIN / "config.json").read_text())["guid"]
        opts = body["config"]["editorConfig"]["plugins"]["options"].get(guid, {}) if status == 200 else {}
        check("editor config: plugin options carry the session token and the page origin", status == 200 and opts.get("appOrigin") == url and opts.get("token") == body["live"]["token"], body.get("live"))
        status, body = api.call("GET", f"/api/office/docs/{docx['id']}/editor?origin=http://evil.example:{url.rsplit(':', 1)[1]}")
        check("editor config refuses a foreign origin (400)", status == 400, str(body))

        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            page = browser.new_page(viewport={"width": 1500, "height": 1000})
            errors: list[str] = []
            page.on("pageerror", lambda e: errors.append(str(e)[:200]))

            # -------------------------------------------------------------- word: every command
            ready, state = open_doc(page, url, docx["id"])
            check("docx opens", ready.startswith("Editor ready"), ready)
            check("plugin connects (word editor)", state.startswith("Assistant live editing: connected (word"), state)
            canvas = editor_canvas(page)
            r = live(page, "get_structure")
            items = r.get("data", {}).get("items", [])
            check("get_structure: the uploaded heading and paragraphs with styles", r["ok"] and items[0]["text"] == "Quarterly report" and items[0].get("heading") == 1 and len(items) == 4, r)
            r = live(page, "insert_text", {"text": "Added at the end.\nSecond added line.", "where": "end"})
            check("insert_text end: two new paragraphs", r["ok"] and r["data"]["paragraphs_added"] == 2, r)
            r = live(page, "insert_paragraphs", {"paragraphs": [{"text": "Executive summary", "style": "Heading 2"}, {"text": "Summary body text."}], "index": 1})
            check("insert_paragraphs at index 1 with a heading style", r["ok"] and r["data"]["at"] == 1 and r["data"]["inserted"] == 2, r)
            r = live(page, "insert_paragraphs", {"paragraphs": [{"text": "x", "style": "No Such Style"}]})
            check("insert_paragraphs with an unknown style is refused (nothing inserted)", not r["ok"] and "unknown paragraph style" in r["error"], r)
            r = live(page, "replace_text", {"search": "Draft", "replace": "Final", "match_case": True})
            check("replace_text: count of replacements", r["ok"] and r["data"]["count"] == 2 and r["data"]["remaining"] == 0, r)
            r = live(page, "get_structure")
            items = r["data"]["items"]
            texts = [i["text"] for i in items]
            idx_costs = texts.index("Final line two about costs.") if "Final line two about costs." in texts else -1
            check("structure after edits: summary heading at 1, replaced text present", texts[1] == "Executive summary" and items[1].get("heading") == 2 and idx_costs > 0, texts)
            r = live(page, "set_paragraph_style", {"from": idx_costs, "style": "Heading 3"})
            check("set_paragraph_style -> Heading 3", r["ok"] and r["data"]["styled"] == 1, r)
            r = live(page, "set_paragraph_style", {"from": 0, "style": "Bogus"})
            check("set_paragraph_style with an unknown style is refused", not r["ok"] and "unknown paragraph style" in r["error"], r)
            r = live(page, "add_comment", {"text": "Check these numbers", "quote": "revenue"})
            check("add_comment on a quote", r["ok"] and r["data"]["on"] == "quote", r)
            r = live(page, "add_comment", {"text": "x", "quote": "not in the document"})
            check("add_comment on missing text is refused", not r["ok"] and "not found" in r["error"], r)
            idx_obsolete = texts.index("Obsolete paragraph to delete.")
            r = live(page, "delete_range", {"from": idx_obsolete, "to": idx_obsolete})
            check("delete_range returns the removed text", r["ok"] and r["data"]["removed_text"] == "Obsolete paragraph to delete.", r)
            r = live(page, "delete_range", {"from": 99})
            check("delete_range out of range is refused", not r["ok"] and "out of range" in r["error"], r)

            # In the editor itself: the user's selection and cursor, and Ctrl+Z undoes a plugin edit.
            box = canvas.bounding_box()
            page.mouse.click(box["x"] + box["width"] / 2, box["y"] + 150)
            page.keyboard.press("Control+a")
            time.sleep(0.5)
            r = live(page, "get_selection")
            check("get_selection reads what the user selected (Ctrl+A)", r["ok"] and "Executive summary" in r["data"]["text"] and "Second added line." in r["data"]["text"], r)
            page.keyboard.press("Control+End")
            time.sleep(0.3)
            r = live(page, "insert_text", {"text": " CURSOR-MARK", "where": "cursor"})
            text = live(page, "get_text")["data"]["text"]
            check("insert_text at the cursor (end of the last paragraph)", r["ok"] and text.endswith("Second added line. CURSOR-MARK"), text[-80:])
            page.keyboard.press("Control+z")
            time.sleep(1)
            text = live(page, "get_text")["data"]["text"]
            check("Ctrl+Z in the editor undoes the plugin's edit", "CURSOR-MARK" not in text and text.endswith("Second added line."), text[-80:])
            page.screenshot(path=str(SHOTS / "docx-after-commands.png"))
            r = live(page, "cells_read")
            check("a spreadsheet command in the word editor is refused with the list", not r["ok"] and "not available in the word editor" in r["error"], r)

            ok, detail = force_save(api, docx["id"])
            check("force-save stores a new version", ok, detail)
            saved = py_docx(READ_DOCX, download(api, docx["id"]))["paragraphs"]
            pairs = [(p["text"], p["style"]) for p in saved if p["text"]]
            expected = [
                ("Quarterly report", "Heading 1"),
                ("Executive summary", "Heading 2"),
                ("Summary body text.", None),
                ("Final line one about revenue.", None),
                ("Final line two about costs.", "Heading 3"),
                ("Added at the end.", None),
                ("Second added line.", None),
            ]
            match = [t for t, _ in pairs] == [t for t, _ in expected] and all(s is None or s == ps for (_, ps), (_, s) in zip(pairs, expected))
            check("saved docx (python-docx): every edit, in order, with the heading styles", match, pairs)
            comments = zip_text(download(api, docx["id"]), "word/comments.xml") if "word/comments.xml" in zip_names(download(api, docx["id"])) else ""
            check("saved docx has the review comment", "Check these numbers" in comments, comments[:200])

            # -------------------------------------------------------------- assistant: reads and edits the open document
            bot = Assistant(page, mock_url)
            bot.open()
            bot.mode("None")
            res = bot.run("CALL:page_snapshot:{}")
            snap = (bot.results("page_snapshot") or [{}])[0]
            check("page_snapshot reports the live document as connected", res == "ran" and snap.get("liveDocument", {}).get("state") == "connected" and snap["liveDocument"].get("editor") == "word", snap.get("liveDocument"))
            tools = bot.sent_tools()
            check("word tools are offered, spreadsheet/slide tools are not", "doc_live_insert_paragraphs" in tools and "doc_live_delete" in tools and "doc_live_cells_write" not in tools and "doc_live_add_slide" not in tools, [t for t in tools if t.startswith("doc_live")])
            res = bot.run('CALL:doc_live_read:{"what":"structure"}\nCALL:doc_live_insert_paragraphs:{"paragraphs":[{"text":"Written by the assistant","style":"Heading 2"},{"text":"Assistant body line."}]}')
            read = (bot.results("doc_live_read") or [{}])[0]
            ins = (bot.results("doc_live_insert_paragraphs") or [{}])[0]
            check("assistant: doc_live_read structure", res == "ran" and read.get("ok") and read["items"][0]["text"] == "Quarterly report", read)
            check("assistant: doc_live_insert_paragraphs edits the open document", ins.get("ok") and ins.get("inserted") == 2, ins)
            text = live(page, "get_text")["data"]["text"]
            check("the assistant's text is in the editor", text.endswith("Written by the assistant\nAssistant body line."), text[-80:])
            res = bot.run("CALL:doc_live_save:{}")
            save = (bot.results("doc_live_save") or [{}])[0]
            check("assistant: doc_live_save writes a new library version", res == "ran" and save.get("saved") is True, save)
            saved = py_docx(READ_DOCX, download(api, docx["id"]))["paragraphs"]
            check("saved file has the assistant's heading", {"text": "Written by the assistant", "style": "Heading 2"} in saved, saved[-3:])
            res = bot.run('CALL:doc_live_cells_write:{"start":"A1","values":[[1]]}')
            wrong = (bot.results("doc_live_cells_write") or [{}])[0]
            check("assistant: a spreadsheet tool on a docx is refused", not wrong.get("ok") and "works in the cell editor" in wrong.get("error", ""), wrong)

            # Trust modes: destructive live edits ask in Risky (Deny keeps the document), plain edits ask in Most.
            bot.mode("Risky")
            before = live(page, "get_text")["data"]["text"]
            res = bot.run('CALL:doc_live_delete:{"from":0,"to":1}')
            check("Risky: doc_live_delete asks as destructive", res.startswith("asked!:Delete paragraphs"), res)
            check("denied delete left the document unchanged", live(page, "get_text")["data"]["text"] == before)
            res = bot.run('CALL:doc_live_replace:{"search":"Final","replace":"Approved"}', decide="Allow")
            rep = (bot.results("doc_live_replace") or [{}])[0]
            check("Risky: doc_live_replace asks as destructive; allowed it replaces", res.startswith("asked!:Replace every") and rep.get("count") == 2, f"{res} {rep}")
            res = bot.run('CALL:doc_live_insert_text:{"text":"Risky mode line."}')
            check("Risky: a plain insert runs without asking", res == "ran" and live(page, "get_text")["data"]["text"].endswith("Risky mode line."), res)
            bot.mode("Most")
            res = bot.run('CALL:doc_live_insert_text:{"text":"never written"}')
            check("Most: an insert asks (not destructive); denied = not written", res.startswith("asked:Edit the open document") and "never written" not in live(page, "get_text")["data"]["text"], res)
            res = bot.run('CALL:doc_live_read:{}')
            check("Most: reading never asks", res == "ran")
            bot.mode("None")
            page.screenshot(path=str(SHOTS / "docx-assistant.png"))

            # Closing the editor: the tools leave the request and a stray call explains what to do.
            key_before_close = api.doc(docx["id"])["key"]
            # Close right after an edit: the editor holds a single user's changes ~2 s before sending; Close waits for them.
            live(page, "insert_text", {"text": "Last-second line.", "where": "end"})
            close_editor(page)
            res = bot.run('CALL:doc_live_read:{}')
            stray = (bot.results("doc_live_read") or [{}])[0]
            check("without an editor the doc_live tools are not offered", not [t for t in bot.sent_tools() if t.startswith("doc_live")])
            check("a stray doc_live call says how to connect", not stray.get("ok") and "No office document is open" in stray.get("error", ""), stray)
            # The Document Server saves a closed session (status 2, new key) a few seconds after the last editor left.
            t_close = time.monotonic()
            rotated = wait_for(lambda: api.doc(docx["id"])["key"] != key_before_close, timeout=120)
            final = py_docx(READ_DOCX, download(api, docx["id"]))["paragraphs"]
            listing = [(d["name"], d["version"]) for d in api.call("GET", "/api/office/docs")[1]["docs"]]
            check("closing saved the trust-mode edits and an edit made just before Close", any(p["text"] == "Approved line one about revenue." for p in final) and any(p["text"] == "Risky mode line." for p in final) and final[-1]["text"] == "Last-second line.",
                  {"rotated": bool(rotated), "after_s": round(time.monotonic() - t_close, 1), "paragraphs": [p["text"] for p in final], "library": listing})

            # -------------------------------------------------------------- word: formatting, tables, lists, header, tracking
            status, body = api.call("POST", "/api/office/docs", {"action": "upload", "name": "Live format.docx", "data": py_docx(MAKE_DOCX)["b64"]})
            fmt_doc = body["doc"]
            ready, state = open_doc(page, url, fmt_doc["id"])
            check("second docx opens and connects", state.startswith("Assistant live editing: connected (word"), state)
            r = live(page, "format_text", {"quote": "revenue", "bold": True, "color": {"r": 192, "g": 0, "b": 0}, "highlight": "yellow"})
            check("format_text on a quote: bold, colour, highlight", r["ok"] and r["data"]["formatted"] == 1, r)
            r = live(page, "format_text", {"from": 0, "italic": True, "size": 20, "font": "Arial"})
            check("format_text on a paragraph: italic, size, font", r["ok"] and r["data"]["formatted"] == 1, r)
            r = live(page, "format_text", {"quote": "nowhere to be found", "bold": True})
            check("format_text on missing text is refused", not r["ok"] and "not found" in r["error"], r)
            r = live(page, "insert_table", {"rows": [["Region", "Sales"], ["North", "120"], ["South", "95"]], "header": True})
            check("insert_table: 3x2 at the end", r["ok"] and r["data"]["rows"] == 3 and r["data"]["cols"] == 2, r)
            r = live(page, "insert_list", {"items": ["First bullet", "Second bullet"]})
            check("insert_list: bullets", r["ok"] and r["data"]["inserted"] == 2, r)
            r = live(page, "insert_list", {"items": ["Step one", "Step two"], "numbered": True})
            check("insert_list: numbered", r["ok"] and r["data"]["numbered"] is True, r)
            r = live(page, "set_header_footer", {"header": "Knife header text", "footer": "Knife footer text"})
            check("set_header_footer", r["ok"] and r["data"].get("header") == "Knife header text", r)
            r = live(page, "track_changes", {"on": True})
            check("track_changes on", r["ok"] and r["data"]["tracking"] is True, r)
            r = live(page, "insert_text", {"text": "Tracked addition.", "where": "end"})
            r = live(page, "get_structure")
            types = [i["type"] for i in r["data"]["items"]]
            check("structure: the table is one element", "table" in types, types)
            page.screenshot(path=str(SHOTS / "docx-format.png"))
            ok, detail = force_save(api, fmt_doc["id"])
            check("formatted docx force-saved", ok, detail)
            raw = download(api, fmt_doc["id"])
            names = zip_names(raw)
            body_xml = zip_text(raw, "word/document.xml")
            headers = " ".join(zip_text(raw, n) for n in names if re.match(r"word/header\d*\.xml$", n))
            footers = " ".join(zip_text(raw, n) for n in names if re.match(r"word/footer\d*\.xml$", n))
            settings = zip_text(raw, "word/settings.xml")
            check("saved docx: bold, red and highlighted run", re.search(r"<w:b/>", body_xml) is not None and re.search(r'<w:color w:val="C00000"', body_xml, re.I) is not None and 'w:highlight w:val="yellow"' in body_xml, body_xml[:300])
            check("saved docx: the table with its cells", "<w:tbl>" in body_xml and "North" in body_xml and "Sales" in body_xml)
            check("saved docx: list paragraphs are numbered", body_xml.count("<w:numPr>") >= 4, body_xml.count("<w:numPr>"))
            check("saved docx: header and footer text", "Knife header text" in headers and "Knife footer text" in footers, (headers[:200], footers[:200]))
            check("saved docx: track changes on, the addition is a revision", "<w:trackRevisions" in settings and "<w:ins " in body_xml and "Tracked addition." in body_xml, settings[:200])
            close_editor(page)

            # -------------------------------------------------------------- spreadsheet
            status, body = api.call("POST", "/api/office/docs", {"action": "create", "format": "xlsx", "name": "Live sheet"})
            xlsx = body["doc"]
            ready, state = open_doc(page, url, xlsx["id"])
            check("plugin connects in the spreadsheet editor", state.startswith("Assistant live editing: connected (cell"), state)
            r = live(page, "cells_write", {"start": "A1", "values": [["Item", "Qty"], ["apples", 3], ["pears", 4], ["total", "=SUM(B2:B3)"]]})
            check("cells_write: a block with a formula", r["ok"] and r["data"]["written"] == 8 and r["data"]["address"] == "A1:B4", r)
            r = live(page, "cells_read", {"range": "A1:B4"})
            check("cells_read: values computed, formula reported", r["ok"] and r["data"]["values"][3] == ["total", "7"] and r["data"]["formulas"] == [{"cell": "B4", "formula": "=SUM(B2:B3)"}], r)
            r = live(page, "cells_write", {"start": "B2", "values": [[10]]})
            check("cells_write refuses to overwrite a filled cell without overwrite", r["ok"] and r["data"].get("needs_overwrite") and r["data"]["non_empty"] == ["B2=3"], r)
            r = live(page, "cells_write", {"start": "B2", "values": [[10]], "overwrite": True})
            check("cells_write with overwrite", r["ok"] and r["data"]["overwritten"] == 1, r)
            r = live(page, "add_sheet", {"name": "Summary"})
            check("add_sheet", r["ok"] and r["data"]["sheets"] == ["Sheet1", "Summary"], r)
            r = live(page, "add_sheet", {"name": "Summary"})
            check("add_sheet with an existing name is refused", not r["ok"], r)
            r = live(page, "get_text")
            check("get_text of the active sheet (the new, empty one)", r["ok"] and r["data"]["sheet"] == "Summary", r)
            r = live(page, "cells_read", {"sheet": "Sheet1"})
            check("cells_read on another sheet by name", r["ok"] and r["data"]["values"][1] == ["apples", "10"], r)
            bot.open() if not page.locator("[role=dialog][aria-label='Toolbox assistant']").count() else None
            bot.run("CALL:page_snapshot:{}")
            tools = bot.sent_tools()
            check("spreadsheet tools offered in the cell editor, word tools not", "doc_live_cells_write" in tools and "doc_live_insert_text" not in tools, [t for t in tools if t.startswith("doc_live")])
            ok, detail = force_save(api, xlsx["id"])
            check("xlsx force-saved", ok, detail)
            raw = download(api, xlsx["id"])
            names = zip_names(raw)
            sheets = [n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml$", n)]
            first = zip_text(raw, sorted(sheets)[0])
            shared = zip_text(raw, "xl/sharedStrings.xml") if "xl/sharedStrings.xml" in names else ""
            workbook = zip_text(raw, "xl/workbook.xml")
            check("saved xlsx: formula, numbers and labels", "<f>SUM(B2:B3)</f>" in first and "apples" in shared and re.search(r"<v>10</v>", first) is not None, first[:400])
            check("saved xlsx: the new sheet", 'name="Summary"' in workbook and len(sheets) == 2, workbook[:300])
            # Formats, a column width, sorting and a chart on Sheet1.
            r = live(page, "cells_format", {"sheet": "Sheet1", "range": "A1:B1", "bold": True, "fill": {"r": 255, "g": 242, "b": 204}, "align": "center"})
            check("cells_format: header bold, filled, centred", r["ok"] and set(r["data"]["set"]) == {"bold", "fill", "align"}, r)
            r = live(page, "cells_format", {"sheet": "Sheet1", "range": "B2:B4", "number_format": "0.00"})
            check("cells_format: number format", r["ok"] and r["data"]["set"] == ["number_format"], r)
            r = live(page, "column_width", {"sheet": "Sheet1", "columns": "A", "first": 0, "last": 0, "width": 24})
            check("column_width", r["ok"], r)
            r = live(page, "sort_range", {"sheet": "Sheet1", "range": "A2:B3", "key": "B2", "descending": False, "header": False})
            check("sort_range by Qty ascending", r["ok"], r)
            r = live(page, "cells_read", {"sheet": "Sheet1", "range": "A2:B3"})
            check("after sorting: pears (4) before apples (10)", r["ok"] and [row[0] for row in r["data"]["values"]] == ["pears", "apples"], r)
            r = live(page, "add_chart", {"sheet": "Sheet1", "abs": "$A$1:$B$3", "type": "bar", "at_col": 3, "at_row": 0, "width_mm": 120, "height_mm": 70, "title": "Fruit"})
            check("add_chart: a bar chart of A1:B3", r["ok"], r)
            page.screenshot(path=str(SHOTS / "xlsx-format.png"))
            ok, detail = force_save(api, xlsx["id"])
            check("formatted xlsx force-saved", ok, detail)
            raw = download(api, xlsx["id"])
            names = zip_names(raw)
            styles = zip_text(raw, "xl/styles.xml")
            sheet1 = zip_text(raw, sorted(n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml$", n))[0])
            # "0.00" is the built-in number format 2: written as numFmtId="2", not as a formatCode.
            fonts = re.search(r"<fonts[^>]*>.*?</fonts>", styles, re.S)
            xfs = re.search(r"<cellXfs[^>]*>.*?</cellXfs>", styles, re.S)
            check("saved xlsx: the 0.00 number format and a bold font",
                  (re.search(r'numFmtId="2"', xfs.group(0) if xfs else "") is not None or 'formatCode="0.00"' in styles) and re.search(r'<b(?: val="(?:1|true)")?\s*/>', fonts.group(0) if fonts else "") is not None,
                  {"fonts": (fonts.group(0) if fonts else "")[:400], "cellXfs": (xfs.group(0) if xfs else "")[:600]})
            check("saved xlsx: column A widened", re.search(r'<col [^>]*min="1"[^>]*width="2[3-5]', sheet1) is not None, re.findall(r"<col [^>]*>", sheet1)[:3])
            check("saved xlsx: a chart part", any(re.match(r"xl/charts/chart\d+\.xml$", n) for n in names), [n for n in names if "chart" in n or "drawing" in n])
            close_editor(page)

            # -------------------------------------------------------------- presentation
            status, body = api.call("POST", "/api/office/docs", {"action": "create", "format": "pptx", "name": "Live deck"})
            pptx = body["doc"]
            ready, state = open_doc(page, url, pptx["id"])
            check("plugin connects in the presentation editor", state.startswith("Assistant live editing: connected (slide"), state)
            r = live(page, "add_slide", {"title": "Quarterly plan", "body": "First point\nSecond point"})
            check("add_slide at the end", r["ok"] and r["data"]["count"] == 2 and r["data"]["added_at"] == 1, r)
            r = live(page, "add_slide", {"title": "Agenda", "index": 0})
            check("add_slide at index 0", r["ok"] and r["data"]["added_at"] == 0 and r["data"]["count"] == 3, r)
            r = live(page, "get_structure")
            slides = r["data"]["slides"]
            check("get_structure: slide texts in order", r["ok"] and slides[0]["texts"] == ["Agenda"] and slides[2]["texts"] == ["Quarterly plan", "First point\nSecond point"], slides)
            page.screenshot(path=str(SHOTS / "pptx.png"))
            ok, detail = force_save(api, pptx["id"])
            check("pptx force-saved", ok, detail)
            raw = download(api, pptx["id"])
            slide_files = sorted(n for n in zip_names(raw) if re.match(r"ppt/slides/slide\d+\.xml$", n))
            all_text = " | ".join(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", zip_text(raw, n))) for n in slide_files)
            check("saved pptx: three slides with the titles and body", len(slide_files) == 3 and "Agenda" in all_text and "Quarterly plan" in all_text and "Second point" in all_text, all_text)
            added = [zip_text(raw, n) for n in slide_files if "Agenda" in zip_text(raw, n) or "Quarterly plan" in zip_text(raw, n)]
            # A new shape's default style paints text in a light colour (white on the white slide): the plugin sets tx1.
            check("added slide text uses the theme text colour (visible on the slide)", len(added) == 2 and all(re.search(r'<a:schemeClr val="tx1"\s*/?>', x) for x in added),
                  [re.findall(r"<a:rPr[^>]*>.{0,200}", x)[:2] for x in added])
            # Slide text by shape, speaker notes, deleting a slide.
            r = live(page, "get_structure")
            plan = r["data"]["slides"][2]
            body_shape = next((sh["shape"] for sh in plan.get("shapes", []) if sh["text"].startswith("First point")), None)
            check("get_structure lists each slide's text shapes", body_shape is not None, plan)
            r = live(page, "set_slide_text", {"index": 2, "shape": body_shape if body_shape is not None else 1, "text": "Revised point\nAnother point"})
            check("set_slide_text replaces a shape's text (returns the old one)", r["ok"] and r["data"]["replaced"].startswith("First point"), r)
            r = live(page, "set_slide_notes", {"index": 2, "text": "Say this out loud."})
            check("set_slide_notes", r["ok"], r)
            r = live(page, "get_structure")
            check("structure: the new text and the notes", r["ok"] and "Revised point\nAnother point" in r["data"]["slides"][2]["texts"] and r["data"]["slides"][2].get("notes") == "Say this out loud.", r["data"]["slides"][2])
            r = live(page, "delete_slide", {"index": 1})
            check("delete_slide returns the texts and the new count", r["ok"] and r["data"]["count"] == 2, r)
            if not page.locator("[role=dialog][aria-label='Toolbox assistant']").count():
                bot.open()
            bot.mode("Risky")
            res = bot.run('CALL:doc_live_delete_slide:{"index":0}')
            check("Risky: deleting a slide asks (destructive); denied = kept", res.startswith("asked!:Delete slide 1") and live(page, "get_structure")["data"]["count"] == 2, res)
            bot.mode("None")
            ok, detail = force_save(api, pptx["id"])
            check("edited pptx force-saved", ok, detail)
            raw = download(api, pptx["id"])
            names = zip_names(raw)
            slide_files = sorted(n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n))
            all_text = " | ".join(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", zip_text(raw, n))) for n in slide_files)
            notes = " ".join(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", zip_text(raw, n))) for n in names if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n))
            check("saved pptx: two slides, the revised text, no old body", len(slide_files) == 2 and "Revised point" in all_text and "First point" not in all_text, all_text)
            check("saved pptx: the speaker notes", "Say this out loud." in notes, notes[:200])
            close_editor(page)
            browser.close()
            check("no page errors", not errors, errors[:5])

        # ------------------------------------------------------------------ a changed plugin is served without a restart
        marker = server.cwd / "sidecars" / "office" / "plugins" / "knife-bridge" / "CHANGED.txt"
        marker.write_text("changed\n")
        status, body = api.call("POST", "/api/services/office", {"action": "ensure"})
        plugins = json.loads(urllib.request.urlopen(public + "/plugins.json", timeout=10).read())
        new_folders = [p.split("/")[-2] for p in plugins.get("pluginsData", [])]
        check("ensure re-syncs a changed plugin under a new folder (old one removed)", status == 200 and len(new_folders) == 1 and new_folders != folders, f"{folders} -> {new_folders}")
        marker.unlink()

        # ------------------------------------------------------------------ an older container is recreated by the next ensure
        # The upgrade path for an existing install: its container was created from a compose file without the plugin
        # mount (and its state has no composeHash). Recreate exactly that with plain compose, then let the app take over.
        idle = wait_for(lambda: api.call("GET", "/api/services/office")[1]["service"]["openSessions"] == 0, timeout=180, step=2)
        check("all editor sessions closed on the server", bool(idle))
        service_dir = Path(server.data_dir) / "services" / "office"
        compose = (service_dir / "compose.yaml").read_text()
        old_compose = "\n".join(line for line in compose.splitlines() if "sdkjs-plugins" not in line) + "\n"
        (service_dir / "compose.yaml").write_text(old_compose)
        up = subprocess.run(["docker", "compose", "-p", project, "-f", "compose.yaml", "up", "-d"], cwd=service_dir, capture_output=True, text=True, timeout=300)
        healthy = wait_for(lambda: _get(public + "/healthcheck") == "true", timeout=180, step=2)
        cid_before = subprocess.run(["docker", "inspect", "--format", "{{.Id}}", container_names(project)[0]], capture_output=True, text=True).stdout.strip()
        old_plugins = json.loads(_get(public + "/plugins.json") or "{}").get("pluginsData")
        check("simulated old container: running without the plugin", up.returncode == 0 and bool(healthy) and old_plugins == [], f"{up.stderr[-300:]} {old_plugins}")
        state_file = service_dir / "state.json"
        st = json.loads(state_file.read_text())
        st.pop("composeHash", None)
        state_file.write_text(json.dumps(st))
        _, body = api.call("GET", "/api/services/office")
        check("status warns about a container from older settings", any("older settings" in w for w in body["service"]["warnings"]), body["service"]["warnings"])
        status, body = api.call("POST", "/api/services/office", {"action": "ensure"})
        recreated = status == 202 and api.wait_op(body["op"], timeout=600)["status"] == "done"
        cid_after = subprocess.run(["docker", "inspect", "--format", "{{.Id}}", container_names(project)[0]], capture_output=True, text=True).stdout.strip()
        new_plugins = json.loads(_get(public + "/plugins.json") or "{}").get("pluginsData", [])
        check("ensure recreated the old container with the plugin mounted", recreated and cid_after != cid_before and len(new_plugins) == 1, f"{status} {cid_before[:12]} -> {cid_after[:12]} {new_plugins}")
        status, body = api.call("POST", "/api/services/office", {"action": "ensure"})
        check("the next ensure is ready at once (no second recreate)", status == 200 and body.get("ready") is True, str(body))
    finally:
        if args.hold:
            input(f"--hold: server {url}, data {server.data_dir}; press Enter to clean up ")
        _, body = api.call("POST", "/api/services/office", {"action": "remove", "purge": True, "removeImage": args.remove_image})
        if isinstance(body, dict) and "op" in body:
            op = api.wait_op(body["op"], timeout=300)
            check("remove (purge) op done", op["status"] == "done", op.get("error", ""))
        left = container_names(project)
        vols = subprocess.run(["docker", "volume", "ls", "-q", "--filter", f"label=com.docker.compose.project={project}"], capture_output=True, text=True).stdout.split()
        check("no container or volume of the e2e project left", not left and not vols, f"{left} {vols}")
        server.stop()
        mock.terminate()
        mock.wait(timeout=10)
    failed = [r for r in RESULTS if not r[1]]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} checks passed; screenshots in {SHOTS}")
    return 1 if failed else 0


if __name__ == "__main__":
    shutil.rmtree(SHOTS, ignore_errors=True)
    sys.exit(main())
