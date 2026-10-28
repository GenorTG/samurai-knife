#!/usr/bin/env python3
"""cv_docx.py: the DOCX opens in python-docx with the expected paragraphs and styles; ATS has no layout tables or images,
the photo layouts place the photo (round = PNG with alpha, square = the JPEG) and at most one layout table.

Run: KNIFE_PYTHON-venv/bin/python tests/cv_docx_test.py   (needs python-docx; synthetic data only)
The view below has the shape of src/lib/cv/render/view.ts CvView; the TS side owns what prints, the script only lays it out.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

SCRIPT = Path(__file__).resolve().parent.parent / "server" / "py" / "scripts" / "cv_docx.py"
NBSP = " "

VIEW = {
    "lang": "pl",
    "name": "Jan Testowy",
    "headline": "Programista backend",
    "contact": [{"text": "Wrocław, Polska", "href": ""}, {"text": "jan@example.test", "href": "mailto:jan@example.test"}],
    "summary": f"Buduję usługi w{NBSP}Pythonie.",  # a stray NBSP must not survive into the DOCX
    "experience": [
        {"heading": "Starszy programista, Acme", "role": "Starszy programista", "org": "Acme", "dates": "03.2022 — obecnie", "place": "Wrocław", "note": "", "bullets": ["Skróciłem build o 40%.", "", "Wdrożyłem CI."]},
        {"heading": "Beta", "role": "", "org": "Beta", "dates": "", "place": "", "note": "", "bullets": []},
    ],
    "projects": [{"name": "Knife", "href": "", "urlText": "knife.example.test", "summary": "Narzędzie w przeglądarce", "tech": "TypeScript, Next.js", "bullets": ["Działa offline."]}],
    "skills": [{"name": "Języki", "items": "Python, TypeScript"}],
    "education": [{"institution": "Politechnika Testowa", "detail": "Informatyka", "dates": "10.2014 — 06.2019"}],
    "languages": ["Polski — ojczysty", "Angielski — C1"],
    "extras": "Małe, przeglądalne zmiany.",
    "consent": "Wyrażam zgodę na przetwarzanie danych osobowych.",
    "titles": {"summary": "Profil", "experience": "Doświadczenie", "projects": "Wybrane projekty", "skills": "Technologie", "education": "Edukacja", "languages": "Języki", "extras": "Sposób pracy"},
    "warnings": [],
    "isEmpty": False,
}


def run_script(request: dict) -> tuple[dict, Path]:
    tmp = Path(tempfile.mkdtemp(prefix="cvdocx-"))
    (tmp / "request.json").write_text(json.dumps({**request, "out_dir": str(tmp / "out")}), "utf-8")
    subprocess.run([sys.executable, str(SCRIPT), str(tmp / "request.json"), str(tmp / "response.json")], check=False, capture_output=True)
    return json.loads((tmp / "response.json").read_text("utf-8")), tmp / "out"


class CvDocxTest(unittest.TestCase):
    def build(self, theme: str = "ats"):
        response, out = run_script({"view": VIEW, "theme": theme, "filename": "Jan-Testowy-CV-pl.docx"})
        self.assertTrue(response["ok"], response)
        self.assertEqual(response["file"], "Jan-Testowy-CV-pl.docx")
        self.assertIn("wordprocessingml.document", response["mime"])
        return Document(out / response["file"]), out / response["file"]

    def test_paragraphs_and_styles(self):
        doc, _ = self.build()
        got = [(p.style.name, p.text) for p in doc.paragraphs]
        self.assertEqual(got[0], ("Title", "Jan Testowy"))
        self.assertEqual(got[1], ("Subtitle", "Programista backend"))
        self.assertEqual(got[2][1], "Wrocław, Polska  |  jan@example.test")
        headings1 = [t for s, t in got if s == "Heading 1"]
        self.assertEqual(headings1, ["Profil", "Doświadczenie", "Wybrane projekty", "Technologie", "Edukacja", "Języki", "Sposób pracy"])
        headings2 = [t for s, t in got if s == "Heading 2"]
        self.assertEqual(headings2, ["Starszy programista, Acme", "Beta", "Knife", "Politechnika Testowa"])
        bullets = [t for s, t in got if s == "List Bullet"]
        self.assertEqual(bullets, ["Skróciłem build o 40%.", "Wdrożyłem CI.", "Działa offline.", "Polski — ojczysty", "Angielski — C1"])
        self.assertIn(("Normal", "03.2022 — obecnie · Wrocław"), got)
        self.assertTrue(any("Języki: " in t or t.startswith("Języki") for _, t in got))
        self.assertEqual(got[-1], ("Normal", "Wyrażam zgodę na przetwarzanie danych osobowych."))
        self.assertTrue(all(t for _, t in got), "no empty paragraphs")

    def test_section_order_and_own_sections(self):
        """`order` decides the print order; a `custom:<id>` section prints its title, item headings, dates and text."""
        view = {
            **VIEW,
            "order": ["summary", "projects", "custom:oss", "experience", "skills", "education", "languages", "extras"],
            "custom": [{"key": "custom:oss", "title": "Open source", "items": [{"title": "Poprawka w X", "text": "Naprawiony wyciek pamięci.", "href": "", "urlText": "github.com/x/x/pull/1", "date": "2025"}]}],
        }
        for theme in ("ats", "sidebar"):
            response, out = run_script({"view": view, "theme": theme, "filename": "cv.docx"})
            self.assertTrue(response["ok"], response)
            doc = Document(out / response["file"])
            paragraphs = [(p.style.name, p.text) for p in doc.paragraphs]
            for table in doc.tables:
                for cell in table._cells:
                    paragraphs += [(p.style.name, p.text) for p in cell.paragraphs]
            headings1 = [t for s, t in paragraphs if s == "Heading 1"]
            if theme == "ats":
                self.assertEqual(headings1, ["Profil", "Wybrane projekty", "Open source", "Doświadczenie", "Technologie", "Edukacja", "Języki", "Sposób pracy"])
            else:  # sidebar: the short lists stay in the left column, the main column follows the order
                main = [h for h in headings1 if h not in ("Technologie", "Języki", "Edukacja")]
                self.assertEqual(main, ["Profil", "Wybrane projekty", "Open source", "Doświadczenie", "Sposób pracy"])
            self.assertIn(("Heading 2", "Poprawka w X · github.com/x/x/pull/1"), paragraphs)
            self.assertIn(("Normal", "Naprawiony wyciek pamięci."), paragraphs)

    def test_no_layout_tables_no_nbsp(self):
        doc, path = self.build("modern")
        self.assertEqual(len(doc.tables), 0)
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertNotIn("<w:tbl>", xml)
        self.assertNotIn("w:txbxContent", xml)
        self.assertNotIn(NBSP, xml)
        self.assertEqual(doc.core_properties.title, "Jan Testowy - CV")
        self.assertEqual(doc.core_properties.language, "pl-PL")

    def photo(self) -> Path:
        from PIL import Image

        path = Path(tempfile.mkdtemp(prefix="cvdocx-photo-")) / "photo.jpg"
        Image.new("RGB", (400, 400), (90, 120, 160)).save(path, "JPEG")
        return path

    def build_with(self, theme: str, photo: Path | None, shape: str = "round"):
        response, out = run_script({"view": VIEW, "theme": theme, "photo": str(photo) if photo else None, "photo_shape": shape, "filename": "cv.docx"})
        self.assertTrue(response["ok"], response)
        return Document(out / response["file"]), out / response["file"]

    def images(self, path: Path) -> list[str]:
        with zipfile.ZipFile(path) as archive:
            return [n for n in archive.namelist() if n.startswith("word/media/")]

    def all_text(self, doc) -> list[str]:
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for cell in table._cells:
                texts.extend(p.text for p in cell.paragraphs)
        return texts

    def test_ats_never_prints_a_photo(self):
        doc, path = self.build_with("ats", self.photo())
        self.assertEqual(self.images(path), [])
        self.assertEqual(len(doc.tables), 0)

    def test_photo_layouts(self):
        for theme, tables in (("modern", 1), ("sidebar", 1), ("elegant", 0)):
            for shape, ext in (("round", ".png"), ("square", ".jpeg")):
                with self.subTest(theme=theme, shape=shape):
                    doc, path = self.build_with(theme, self.photo(), shape)
                    media = self.images(path)
                    self.assertEqual(len(media), 1, media)
                    self.assertTrue(media[0].endswith(ext) or media[0].endswith(".jpg"), media)
                    self.assertEqual(len(doc.tables), tables)
                    text = self.all_text(doc)
                    for needed in ("Jan Testowy", "Programista backend", "Skróciłem build o 40%.", "Politechnika Testowa", "Wyrażam zgodę na przetwarzanie danych osobowych."):
                        self.assertIn(needed, text)

    def test_photo_layouts_without_photo(self):
        doc, path = self.build_with("modern", None)
        self.assertEqual((len(doc.tables), self.images(path)), (0, []))
        doc, _ = self.build_with("elegant", None)
        self.assertEqual(doc.paragraphs[0].text, "Jan Testowy")
        self.assertEqual(doc.styles["Title"].font.name, "Georgia")
        doc, path = self.build_with("sidebar", None)
        self.assertEqual(len(doc.tables), 1)
        side = [p.text for p in doc.tables[0].rows[0].cells[0].paragraphs]
        main = [p.text for p in doc.tables[0].rows[0].cells[1].paragraphs]
        self.assertIn("jan@example.test", side)
        self.assertIn("Technologie", side)
        self.assertEqual(main[0], "Jan Testowy")
        self.assertIn("Doświadczenie", main)
        self.assertNotIn("Doświadczenie", side)

    def xml_and_rels(self, path: Path) -> tuple[str, str]:
        with zipfile.ZipFile(path) as archive:
            return archive.read("word/document.xml").decode("utf-8"), archive.read("word/_rels/document.xml.rels").decode("utf-8")

    def test_every_layout_builds_with_everything_in_it(self):
        """All 14 layouts (layout.ts CV_LAYOUTS) build; ATS-safe ones keep no tables; the text is all there."""
        view = {**VIEW, "contact": [*VIEW["contact"], {"text": "github.com/jantest", "href": "https://github.com/jantest", "kind": "github"}]}
        for theme in ("ats", "compact", "minimal", "classic", "modern", "elegant", "executive", "technical", "timeline", "banner", "sidebar", "rightbar", "creative", "europass"):
            with self.subTest(theme=theme):
                response, out = run_script({"view": view, "theme": theme, "look": {"accent": "#2f6b4f", "paper": "letter"}, "filename": "cv.docx"})
                self.assertTrue(response["ok"], response)
                doc = Document(out / response["file"])
                text = "\n".join(self.all_text(doc))
                for needed in ("Jan Testowy", "Skróciłem build o 40%.", "Politechnika Testowa", "Angielski", "Wyrażam zgodę"):
                    self.assertIn(needed, text)
                if theme in ("ats", "compact", "minimal", "classic"):
                    self.assertEqual(len(doc.tables), 0)
                xml, rels = self.xml_and_rels(out / response["file"])
                self.assertIn("<w:hyperlink", xml, "contact links are real hyperlinks")
                self.assertIn('Target="https://github.com/jantest"', rels)
                self.assertIn('TargetMode="External"', rels)
                width = doc.sections[0].page_width
                self.assertAlmostEqual(width.inches, 8.5, places=1)

    def test_project_pages(self):
        page = {"name": "Knife", "href": "https://knife.example.test", "urlText": "knife.example.test", "summary": "Narzędzie w przeglądarce", "tech": ["TypeScript"], "role": "Autor", "outcome": "200 użytkowników", "story": ["Pierwszy akapit.", "Drugi akapit."], "bullets": ["Działa offline."]}
        view = {**VIEW, "pages": [page], "pageTitles": {"role": "Moja rola", "outcome": "Efekt", "highlights": "Najważniejsze", "tech": "Technologie", "heading": "Projekt"}}
        response, out = run_script({"view": view, "theme": "modern", "page_mode": "cv+pages", "filename": "cv.docx"})
        self.assertTrue(response["ok"], response)
        doc = Document(out / response["file"])
        text = self.all_text(doc)
        self.assertIn("Jan Testowy", text)
        for needed in ("Pierwszy akapit.", "Drugi akapit.", "Moja rola", "200 użytkowników", "Efekt"):
            self.assertIn(needed, text)
        xml, rels = self.xml_and_rels(out / response["file"])
        self.assertIn('w:type="page"', xml)
        self.assertIn('Target="https://knife.example.test"', rels)
        # Pages only: no CV body, just the showcase.
        response, out = run_script({"view": view, "theme": "ats", "page_mode": "pages", "filename": "p.docx"})
        self.assertTrue(response["ok"], response)
        text = self.all_text(Document(out / response["file"]))
        self.assertIn("Pierwszy akapit.", text)
        self.assertNotIn("Skróciłem build o 40%.", text)
        # Pages only, but no page: a readable refusal.
        response, _ = run_script({"view": VIEW, "page_mode": "pages", "filename": "p.docx"})
        self.assertFalse(response["ok"])

    def test_bad_requests_are_readable(self):
        for request, fragment in (({"view": "x"}, "view"), ({"view": {**VIEW, "lang": "de"}}, "lang"), ({"view": VIEW, "theme": "neon"}, "theme"), ({"view": VIEW, "theme": "modern", "photo": "/nonexistent.jpg"}, "photo"), ({"view": VIEW, "photo_shape": "star"}, "photo_shape"), ({"view": VIEW, "look": {"accent": "red"}}, "look.accent"), ({"view": VIEW, "page_mode": "x"}, "page_mode"), ({"view": {"lang": "en", "titles": {}}}, "nothing to print")):
            response, _ = run_script(request)
            self.assertFalse(response["ok"])
            self.assertIn(fragment, response["error"])


if __name__ == "__main__":
    unittest.main()
